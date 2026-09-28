"""
내역서 물량산출 자동화 웹앱
실행: streamlit run app.py
"""

import streamlit as st
import pandas as pd
from openpyxl.styles import PatternFill
from io import BytesIO

st.set_page_config(page_title="내역서 검토 자동화", layout="wide")

# 실제 내역서의 컬럼명이 다르면 오른쪽 값만 수정
COLUMN_MAP = {
    "공종": "공종",
    "품명": "품명",
    "규격": "규격",
    "단위": "단위",
    "수량": "수량",
    "재료비단가": "재료비단가",
    "노무비단가": "노무비단가",
    "경비단가": "경비단가",
    "합계단가": "합계단가",
    "금액": "금액",
}


def load_naeyoekseo(file, sheet_name=0) -> pd.DataFrame:
    df = pd.read_excel(file, sheet_name=sheet_name)
    df = df.rename(columns={v: k for k, v in COLUMN_MAP.items()})
    df = df.dropna(subset=["품명"])
    return df


def validate_amount(df: pd.DataFrame) -> pd.DataFrame:
    df["계산금액"] = df["수량"] * df["합계단가"]
    df["오류여부"] = (df["계산금액"] - df["금액"]).abs() > 1
    return df


def summarize_by_gongjong(df: pd.DataFrame) -> pd.DataFrame:
    summary = (
        df.groupby("공종")
        .agg(항목수=("품명", "count"), 총수량=("수량", "sum"), 총금액=("금액", "sum"))
        .reset_index()
    )
    summary["금액비중(%)"] = (summary["총금액"] / summary["총금액"].sum() * 100).round(2)
    return summary.sort_values("총금액", ascending=False)


def build_report_bytes(df: pd.DataFrame, summary: pd.DataFrame) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="공종별집계", index=False)
        df.to_excel(writer, sheet_name="전체검토", index=False)
        ws = writer.sheets["전체검토"]
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        error_col_idx = df.columns.get_loc("오류여부") + 1
        for row in range(2, len(df) + 2):
            if ws.cell(row=row, column=error_col_idx).value:
                for col in range(1, len(df.columns) + 1):
                    ws.cell(row=row, column=col).fill = red_fill
    return output.getvalue()


st.title("📋 내역서 물량산출 자동화")
st.write("내역서 엑셀 파일을 업로드하면 공종별 집계와 금액 오류를 자동으로 검토합니다.")

uploaded_file = st.file_uploader("내역서 엑셀 파일 업로드 (.xlsx)", type=["xlsx"])

if uploaded_file is not None:
    try:
        df = load_naeyoekseo(uploaded_file)
        df = validate_amount(df)
        summary = summarize_by_gongjong(df)

        error_count = int(df["오류여부"].sum())
        total_amount = int(df["금액"].sum())

        c1, c2, c3 = st.columns(3)
        c1.metric("전체 항목 수", f"{len(df):,}개")
        c2.metric("오류 항목 수", f"{error_count:,}개")
        c3.metric("총 금액", f"{total_amount:,}원")

        st.subheader("공종별 집계")
        st.dataframe(summary, use_container_width=True)

        st.subheader("전체 항목 검토")
        if error_count > 0:
            st.warning(f"금액 계산 오류가 {error_count}건 발견되었습니다.")
        st.dataframe(
            df.style.apply(
                lambda row: ["background-color: #ffc7ce" if row["오류여부"] else "" for _ in row],
                axis=1,
            ),
            use_container_width=True,
        )

        st.download_button(
            label="검토결과 리포트 다운로드 (.xlsx)",
            data=build_report_bytes(df, summary),
            file_name="검토결과_리포트.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    except KeyError as e:
        st.error(f"컬럼을 찾을 수 없습니다: {e}. COLUMN_MAP을 확인하세요.")
    except Exception as e:
        st.error(f"파일 처리 중 오류가 발생했습니다: {e}")
else:
    st.info("파일을 업로드하면 결과가 여기에 표시됩니다.")
